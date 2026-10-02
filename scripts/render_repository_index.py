"""Generate the development index from the canonical repository relationship graph."""
from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import subprocess

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from render_repository_map import ROOT, collect


def generate(output: Path) -> tuple[int, int]:
    graph = collect()
    paths, edges, mapping = graph['paths'], graph['edges'], graph['mapping']
    tracked = set(subprocess.check_output(['git', 'ls-files', '-z'], cwd=ROOT).decode('utf-8').split('\0'))
    areas = {a['path']: a for a in mapping['areas']}
    counts = Counter(p.split('/')[0] if '/' in p else 'raiz' for p in paths)
    outgoing, incoming = {}, {}
    for source, target, relation in edges:
        outgoing.setdefault(source, set()).add(target)
        incoming.setdefault(target, set()).add(source)
    wb = Workbook()
    guide = wb.active
    guide.title = 'Orientação'
    files = wb.create_sheet('Arquivos')
    relations = wb.create_sheet('Relações')

    def put(sheet, row, values):
        for col, value in enumerate(values, 1):
            sheet.cell(row, col, value)

    def link(cell, p, label, remote=False):
        from urllib.parse import quote
        cell.value = label
        cell.hyperlink = ('https://github.com/duducel204/AutoCompiler/blob/main/' + quote(p, safe='/')) if remote else (ROOT / p).as_uri()
        cell.font = Font(name='Arial', size=10, color='0563C1', underline='single')

    def title(p):
        import re
        if p.endswith(('.md', '.py')):
            text = (ROOT / p).read_text(encoding='utf-8-sig')
            match = re.search(r'^#\s+(.+)$', text, re.M) if p.endswith('.md') else re.search(r'^\s*(?:"""|\x27\x27\x27)([^\n]+)', text)
            if match:
                return match[1].replace('"""', '').replace("'''", '').strip()[:180]
        area = p.split('/')[0]
        return {'src':'Implementação: consultar imports e fluxos', 'tests':'Contrato ou cenário de teste', 'data':'Registro estruturado', 'schemas':'Contrato de dados', 'examples':'Entrada de exemplo', 'generated':'Saída compilada', 'scripts':'Comando reutilizável', 'web':'Componente da interface', '.github':'Workflow de integração', 'skills':'Procedimento reutilizável', 'docs':'Documento ou índice'}.get(area, 'Configuração ou entrada da raiz')

    put(files, 2, ['Índice principal de desenvolvimento'])
    put(files, 3, ['Filtre por área, tipo ou fluxo. Consulte Relações antes de alterar.'])
    put(files, 4, ['Links locais dependem da máquina. GitHub aponta à main; arquivos novos aguardam publicação.'])
    put(files, 6, ['Caminho', 'Área', 'Tipo', 'Função / título', 'Fluxos', 'Dependências', 'Referências recebidas', 'GitHub main', 'Arquivo local', 'Orientação'])
    for row, p in enumerate(paths, 7):
        area = p.split('/')[0] if '/' in p else 'raiz'
        flows = ', '.join(f['id'] for f in mapping['flows'] if p in f['stages'] + f['docs'])
        out, inc = len(outgoing.get(p, [])), len(incoming.get(p, []))
        put(files, row, [p, area, Path(p).suffix[1:] or 'sem extensão', title(p), flows or 'Sem fluxo editorial', out, inc, 'Ainda não versionado', 'Abrir localmente', 'Consultar Relações' if out+inc else 'Sem relação extraída; revisar contexto'])
        if p in tracked:
            link(files.cell(row, 8), p, 'Abrir no GitHub', True)
        link(files.cell(row, 9), p, 'Abrir localmente')
    put(relations, 2, ['Relações entre arquivos'])
    put(relations, 3, ['Filtre por Origem ou Destino. Imports indicam dependências, não execução.'])
    put(relations, 6, ['Origem', 'Relação', 'Destino', 'Link origem', 'Link destino'])
    for row, (source, target, relation) in enumerate(edges, 7):
        put(relations, row, [source, relation, target, 'Abrir origem', 'Abrir destino'])
        link(relations.cell(row, 4), source, 'Abrir origem')
        link(relations.cell(row, 5), target, 'Abrir destino')

    put(guide, 2, ['AutoCompiler: orientação de desenvolvimento'])
    put(guide, 4, ['Comece aqui', 'Objetivo', 'Referência'])
    for row, label, objective, p in [(5,'Entender o projeto','Propósito, regras e estado','docs/INDEX.md'), (6,'Escolher onde alterar','Áreas e fluxos','docs/REPOSITORY_MAP.md'), (7,'Delimitar o trabalho','Limites e próximos passos','docs/NEXT_TASK_PREP.md'), (8,'Executar com consistência','Contribuição e evidências','CONTRIBUTING.md')]:
        put(guide, row, [label, objective, 'Abrir documento'])
        link(guide.cell(row, 3), p, 'Abrir documento', True)
    put(guide, 10, ['Arquivos por área'])
    put(guide, 12, ['Área', 'Arquivos', 'Responsabilidade'])
    for row, area in enumerate(['raiz'] + list(areas), 13):
        put(guide, row, [area, counts[area], areas.get(area, {}).get('role', 'Entradas e políticas do repositório')])
    put(guide, 25, ['Total de arquivos', len(paths)])
    guide.cell(10, 5, 'Fluxos registrados')
    for col, value in enumerate(['ID', 'Fluxo', 'Entrada', 'Saída'], 5):
        guide.cell(12, col, value)
    for row, flow in enumerate(mapping['flows'], 13):
        for col, value in enumerate([flow['id'], flow['name'], flow['input'], flow['output']], 5):
            guide.cell(row, col, value)
    put(guide, 28, ['Fontes e atualização'])
    for row, text in enumerate([
        'Fonte: inventário Git e grafo de render_repository_map.py; fluxos em data/repository_map.json.',
        'Atualizar: python scripts/render_repository_index.py (requer openpyxl).',
        'Snapshot gerado das fontes: totais são valores, atualizados na regeneração.',
        'Inclui arquivos versionados e novos nas áreas do projeto; exclui estado ignorado e arquivos pessoais não versionados na raiz.',
        'As relações orientam leitura; não comprovam execução nem promovem capacidades.'
    ], 29):
        put(guide, row, [text])

    for sheet, widths, headers in [(guide,[30,27,74,3,22,46,62,62],[4,12]), (files,[65,12,14,68,28,17,22,24,24,45],[6]), (relations,[68,46,68,24,24],[6])]:
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            sheet.row_dimensions[row[0].row].height = 48 if sheet != guide and row[0].row > 6 else 30
            for cell in row:
                if not cell.hyperlink:
                    cell.font = Font(name='Arial', size=10, color='172B4D')
                cell.alignment = Alignment(vertical='center', wrap_text=cell.row > 6)
        for col, width in enumerate(widths, 1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet['A2'].font = Font(name='Arial', size=16, bold=True, color='172B4D')
        sheet['A2'].alignment = Alignment(wrap_text=False)
        for r in headers:
            for cell in sheet[r]:
                if cell.value is not None:
                    cell.fill = PatternFill('solid', fgColor='263B56')
                    cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r in (3,4) if sheet != guide else range(29,34):
            sheet.cell(r,1).alignment = Alignment(wrap_text=False, vertical='center')
    for row in range(13, 19):
        guide.row_dimensions[row].height = 78
    files.freeze_panes = 'B7'
    relations.freeze_panes = 'A7'
    for sheet, ref, name in [(files,f'A6:J{len(paths)+6}','RepositoryFiles'),(relations,f'A6:E{len(edges)+6}','RepositoryRelations')]:
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    saved = load_workbook(output)
    actual = [saved['Arquivos'].cell(r,1).value for r in range(7,len(paths)+7)]
    assert actual == paths and len(set(actual)) == len(paths)
    assert saved['Orientação']['B25'].value == sum(counts.values())
    assert all(saved['Arquivos'].cell(r,9).hyperlink for r in range(7,len(paths)+7))
    assert all(saved['Relações'].cell(r,c).hyperlink for r in range(7,len(edges)+7) for c in (4,5))
    print(f'Generated {output}: {len(paths)} files; {len(edges)} relationships. Saved inventory, totals and hyperlinks verified.')
    return len(paths), len(edges)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=ROOT / 'docs/REPOSITORY_INDEX.xlsx')
    args = parser.parse_args()
    generate(args.output.resolve())
